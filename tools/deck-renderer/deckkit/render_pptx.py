"""Draw ops -> editable .pptx (opens natively in Keynote and PowerPoint)."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.parts.image import Image
from pptx.util import Emu, Pt

from . import core

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def rgb(hex6):
    return RGBColor.from_string(hex6.lstrip("#").upper())


def build(deck, ops_per_slide, theme, out_path):
    prs = Presentation()
    prs.slide_width = Emu(int(Pt(core.W)))
    prs.slide_height = Emu(int(Pt(core.H)))
    blank = prs.slide_layouts[6]
    for slide_data, ops in zip(deck["slides"], ops_per_slide):
        s = prs.slides.add_slide(blank)
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(theme["bg"])
        for op in ops:
            _emit(s.shapes, op, theme)
        if slide_data.get("notes"):
            s.notes_slide.notes_text_frame.text = slide_data["notes"]
    _set_core_properties(prs, deck["meta"])
    prs.save(out_path)
    return out_path


def _set_core_properties(prs, meta):
    cp = prs.core_properties
    cp.title = str(meta.get("title", ""))
    if meta.get("author"):
        cp.author = str(meta["author"])
    if meta.get("subtitle"):
        cp.subject = str(meta["subtitle"])


def _emit(shapes, op, th):
    kind = op["op"]
    if kind == "rect":
        _rect(shapes, op)
    elif kind == "text":
        _text(shapes, op, th)
    elif kind == "poly":
        _poly(shapes, op)
    elif kind == "image":
        _image(shapes, op)


def _rect(shapes, op):
    w, h = max(op["w"], 0.1), max(op["h"], 0.1)
    radius = op.get("radius", 0)
    if radius > 0 and min(w, h) > 1.5:
        shp = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(op["x"]), Pt(op["y"]), Pt(w), Pt(h))
        try:
            shp.adjustments[0] = min(0.5, radius / min(w, h))
        except (IndexError, ValueError):
            pass
    else:
        shp = shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(op["x"]), Pt(op["y"]), Pt(w), Pt(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(op["fill"])
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _text(shapes, op, th):
    box = shapes.add_textbox(Pt(op["x"]), Pt(op["y"]), Pt(max(op["w"], 1)), Pt(max(op["h"], 1)))
    tf = box.text_frame
    tf.word_wrap = bool(op.get("wrap", True))
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR.get(op.get("valign", "t"), MSO_ANCHOR.TOP)
    lines = str(op["s"]).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(op.get("align", "l"), PP_ALIGN.LEFT)
        p.line_spacing = op.get("leading", 1.35)
        run = p.add_run()
        run.text = line
        f = run.font
        f.size = Pt(op["size"])
        f.bold = bool(op.get("bold"))
        f.color.rgb = rgb(op["color"])
        _typeface(run, th, op.get("font", "sans"))
        if op.get("tracking"):
            run._r.get_or_add_rPr().set("spc", str(int(round(op["tracking"] * 100))))
    return box


def _typeface(run, th, kind):
    latin, cjk = core.face(th, kind)
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:latin", latin), ("a:ea", cjk), ("a:cs", latin)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", face)


def _poly(shapes, op):
    pts = op["pts"]
    if len(pts) < 2:
        return None
    builder = shapes.build_freeform(pts[0][0], pts[0][1], Pt(1))
    builder.add_line_segments([(x, y) for x, y in pts[1:]], close=op.get("close", False))
    shp = builder.convert_to_shape()
    if op.get("fill"):
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(op["fill"])
    else:
        shp.fill.background()
    if op.get("color"):
        shp.line.color.rgb = rgb(op["color"])
        shp.line.width = Pt(op.get("weight", 1.8))
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _image(shapes, op):
    x, y, w, h = op["x"], op["y"], op["w"], op["h"]
    try:
        px, py = Image.from_file(op["path"]).size
        scale = min(w / px, h / py)
        dw, dh = px * scale, py * scale
        x, y = x + (w - dw) / 2, y + (h - dh) / 2
        w, h = dw, dh
    except Exception:
        pass
    return shapes.add_picture(op["path"], Pt(x), Pt(y), Pt(w), Pt(h))
